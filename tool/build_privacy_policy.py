from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE
from pathlib import Path

SOURCE = Path(r"C:\Users\Edwin\Downloads\Politica_de_Privacidad_DashCore.docx")
OUTPUT = Path(r"E:\dashcore\dashcore\dashcore\docs\Politica_de_Privacidad_DashCore_v2.docx")
OUTPUT.parent.mkdir(parents=True, exist_ok=True)

doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Inches(8.5), Inches(11)
sec.top_margin = sec.bottom_margin = sec.left_margin = sec.right_margin = Inches(1)
sec.header_distance = sec.footer_distance = Inches(.492)

navy, blue, cyan, muted, pale = "0B1726", "145DA0", "00AFC8", "5B6573", "EEF5F8"

styles = doc.styles
normal = styles["Normal"]
normal.font.name = "Aptos"
normal.font.size = Pt(10.5)
normal.font.color.rgb = RGBColor.from_string(navy)
normal.paragraph_format.space_after = Pt(6)
normal.paragraph_format.line_spacing = 1.12
for name, size, color, before, after in [
    ("Title", 28, navy, 0, 8), ("Subtitle", 11, muted, 0, 18),
    ("Heading 1", 16, blue, 16, 8), ("Heading 2", 12.5, blue, 11, 5),
    ("Heading 3", 11, navy, 8, 4)]:
    s = styles[name]
    s.font.name, s.font.size = "Aptos Display", Pt(size)
    s.font.color.rgb = RGBColor.from_string(color)
    s.font.bold = name != "Subtitle"
    s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(before), Pt(after)
    s.paragraph_format.keep_with_next = True

if "Callout" not in styles:
    callout = styles.add_style("Callout", WD_STYLE_TYPE.PARAGRAPH)
else:
    callout = styles["Callout"]
callout.font.name, callout.font.size = "Aptos", Pt(10.5)
callout.font.color.rgb = RGBColor.from_string(navy)
callout.paragraph_format.left_indent = Inches(.18)
callout.paragraph_format.right_indent = Inches(.18)
callout.paragraph_format.space_before = callout.paragraph_format.space_after = Pt(8)

def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd'); shd.set(qn('w:fill'), fill); tcPr.append(shd)

def border_bottom(paragraph, color=cyan, size="18"):
    pPr = paragraph._p.get_or_add_pPr(); pBdr = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom'); bottom.set(qn('w:val'), 'single')
    bottom.set(qn('w:sz'), size); bottom.set(qn('w:space'), '8'); bottom.set(qn('w:color'), color)
    pBdr.append(bottom); pPr.append(pBdr)

def add_bullet(text):
    p = doc.add_paragraph(style="List Bullet"); p.add_run(text); return p

def add_section(title, paragraphs=(), bullets=()):
    doc.add_heading(title, level=1)
    for text in paragraphs: doc.add_paragraph(text)
    for text in bullets: add_bullet(text)

header = sec.header.paragraphs[0]
header.text = "DASHCORE  /  PRIVACIDAD"
header.style = styles["Caption"]
header.runs[0].font.color.rgb = RGBColor.from_string(muted)
header.runs[0].font.bold = True
border_bottom(header, "D8E4EA", "6")

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
run = footer.add_run("DashCore · Política de Privacidad v2.0 · 4 de septiembre de 2026")
run.font.size, run.font.color.rgb = Pt(8), RGBColor.from_string(muted)

p = doc.add_paragraph(style="Title"); p.add_run("Política de Privacidad")
border_bottom(p)
p = doc.add_paragraph(style="Subtitle")
p.add_run("DASHCORE  ·  VERSIÓN 2.0  ·  VIGENTE DESDE EL 4 DE SEPTIEMBRE DE 2026")

table = doc.add_table(rows=3, cols=2)
table.alignment = WD_TABLE_ALIGNMENT.CENTER
table.autofit = False
labels = [("Aplicación", "DashCore"), ("Responsable", "[NOMBRE LEGAL DEL RESPONSABLE]"),
          ("Contacto de privacidad", "[privacy@dominio.com]")]
for row, (a,b) in zip(table.rows, labels):
    row.cells[0].width, row.cells[1].width = Inches(1.65), Inches(4.85)
    row.cells[0].text, row.cells[1].text = a, b
    shade(row.cells[0], pale)
    row.cells[0].paragraphs[0].runs[0].font.bold = True
    for c in row.cells: c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER

p = doc.add_paragraph(style="Callout")
p.add_run("EN POCAS PALABRAS\n").bold = True
p.add_run("DashCore muestra información del vehículo y puede sincronizar una cuenta. La analítica, las mediciones de rendimiento y el envío de GPS o telemetría en vivo están desactivados por defecto y requieren decisiones separadas y revocables.")

add_section("1. Alcance", [
    "Esta política explica qué información trata DashCore, por qué la utiliza, durante cuánto tiempo puede conservarla y qué controles tiene el usuario. Se aplica a la aplicación móvil, autenticación, perfiles, dashboards, OBD2, GPS, mapas, diagnóstico, recursos, actualizaciones y servicios remotos asociados.",
    "Aceptar esta política permite continuar con el flujo de cuenta. La aceptación no sustituye los permisos de Android ni activa por sí sola tratamientos opcionales."
])

add_section("2. Información que tratamos", bullets=[
    "Cuenta: correo, identificador de usuario, nombre o imagen de perfil y metadatos de autenticación.",
    "Dispositivo: fabricante, modelo, versión de Android, arquitectura, versión de DashCore e información técnica necesaria para compatibilidad.",
    "Vehículo: marca, modelo, año, motor y preferencias asociadas.",
    "OBD2: velocidad, RPM, temperatura, voltaje, combustible, odómetro, códigos de diagnóstico y estado de conexión disponibles para la función utilizada.",
    "Conducción: distancia, viajes, tiempo de conducción y velocidad máxima calculados durante una sesión válida.",
    "GPS: coordenadas, precisión, rumbo, velocidad y fecha/hora, únicamente cuando una función autorizada lo requiera.",
    "Soporte y rendimiento: informes enviados voluntariamente y métricas técnicas activadas por el usuario."
])

add_section("3. GPS y telemetría en vivo", [
    "DashCore está preparado para ofrecer un panel remoto que muestre vehículos asociados a una cuenta. Esta función debe permanecer desactivada hasta que el usuario active expresamente “GPS y telemetría en vivo”.",
    "Cuando se active, la aplicación podrá transmitir periódicamente la ubicación precisa, precisión, rumbo, velocidad GPS, estado de conexión OBD y parámetros actuales del vehículo. Estos datos pueden revelar desplazamientos, rutinas y ubicación del usuario, por lo que se consideran especialmente sensibles.",
    "La función debe poder desactivarse desde Ajustes. También puede impedirse revocando el permiso de ubicación en Android. Desactivarla detiene nuevas transmisiones, pero no elimina automáticamente registros históricos previamente autorizados."
], bullets=[
    "Frecuencia prevista para la vista en vivo: intervalo limitado y configurable, no una transmisión sin control.",
    "Acceso remoto: exclusivamente el titular de la cuenta y personal autorizado bajo controles de acceso.",
    "No se habilitarán comandos remotos al vehículo mediante esta vista.",
    "El historial de recorridos, si se añade, tendrá consentimiento y retención separados."
])

add_section("4. Finalidades y bases", [
    "Los datos necesarios para crear la cuenta, mantener la sesión, guardar preferencias y prestar las funciones solicitadas se tratan para ejecutar el servicio. La analítica no esencial, medición de rendimiento y seguimiento en vivo se basan en una elección voluntaria cuando corresponda. DashCore también puede tratar datos mínimos para seguridad, prevención de abuso y cumplimiento legal."
])

add_section("5. Funcionamiento local y ACC", [
    "En radios Android, DashCore puede mantener un servicio visible mientras existe una sesión activa. Al detectar la desconexión de ACC, la aplicación reduce su actividad, detiene el sondeo OBD2, desconecta el adaptador y suspende GPS y tareas no esenciales. Al regresar ACC, intenta restaurar la sesión previamente activa. El fabricante de la radio y Android pueden imponer límites adicionales o finalizar procesos."
])

add_section("6. Consentimientos y controles", bullets=[
    "Política de privacidad: se presenta antes de iniciar sesión o crear una cuenta y se registra la versión aceptada.",
    "Analítica: opción independiente, desactivada por defecto y revocable.",
    "Medición de rendimiento: opción independiente en Ajustes; sólo se envía al backend cuando también existe el consentimiento requerido.",
    "GPS y telemetría en vivo: opción independiente, desactivada por defecto y acompañada del permiso de Android.",
    "Retirar un consentimiento detiene nuevos tratamientos basados en él y no afecta el tratamiento anterior legítimo."
])

add_section("7. Proveedores y transferencias", [
    "DashCore utiliza Supabase para autenticación, base de datos y sincronización. El inicio con Google implica además el tratamiento realizado por Google. Los mapas pueden utilizar proveedores de teselas o datos geográficos. Algunos proveedores pueden tratar datos fuera del país del usuario con las salvaguardas exigibles."
])

add_section("8. Conservación", [
    "Los datos se conservarán sólo durante el tiempo necesario para la cuenta, el servicio, seguridad, soporte y obligaciones legales. La posición en vivo debe reemplazarse por el estado más reciente y eliminarse al vencer un plazo corto. Los eventos y métricas deben tener políticas automáticas de retención. Antes de producción, DashCore debe publicar los plazos definitivos."
])

add_section("9. Seguridad", [
    "DashCore aplica cifrado en tránsito, autenticación y políticas de seguridad por fila para separar los datos de cada usuario. Las claves administrativas nunca deben distribuirse dentro de la aplicación. Ningún sistema ofrece seguridad absoluta; los incidentes se gestionarán conforme a la ley aplicable."
])

add_section("10. Derechos", [
    "Según la jurisdicción, el usuario puede solicitar información, acceso, rectificación, supresión, limitación, portabilidad u oposición, así como retirar consentimientos. Las solicitudes se dirigirán a [privacy@dominio.com]. Puede requerirse verificación razonable de identidad."
])

add_section("11. Menores", [
    "DashCore no está dirigido deliberadamente a menores que no puedan consentir legalmente. Un representante puede solicitar revisión o eliminación mediante el canal de privacidad."
])

add_section("12. Cambios y contacto", [
    "Los cambios materiales se comunicarán antes de aplicar nuevas finalidades cuando corresponda. Una nueva versión podrá requerir aceptación o consentimiento renovado.",
    "Responsable: [NOMBRE LEGAL] · Correo: [privacy@dominio.com] · Domicilio: [DOMICILIO / PAÍS] · Web: [URL DE PRIVACIDAD]"
])

add_section("Anexo operativo antes de publicar", bullets=[
    "Completar nombre legal, correo, domicilio, sitio web y mecanismo de eliminación de cuenta.",
    "Definir retención exacta para ubicación en vivo, sesiones, eventos, informes y consentimientos.",
    "Aplicar RLS en todas las tablas y separar roles de usuario, soporte y administrador.",
    "No activar la transmisión en vivo hasta aprobar el aviso específico dentro de Ajustes.",
    "Realizar revisión jurídica para los países donde se distribuirá DashCore."
])

doc.add_heading("Referencias regulatorias", level=1)
for item in [
    "Reglamento (UE) 2016/679 (RGPD).",
    "California Consumer Privacy Act, cuando resulte aplicable.",
    "República Dominicana: Ley No. 172-13 sobre protección de datos personales.",
    "Política de Datos de Usuario de Google Play y requisitos de divulgación destacada."]:
    add_bullet(item)

doc.core_properties.title = "Política de Privacidad de DashCore"
doc.core_properties.subject = "Privacidad, telemetría OBD2, GPS y servicios en vivo"
doc.core_properties.author = "DashCore"
doc.core_properties.comments = "Documento sujeto a revisión jurídica antes de publicación."
doc.save(OUTPUT)
print(OUTPUT)
